"""
Create or overwrite an Excel workbook using openpyxl.
This script is invoked from the Electron main process to generate a table-only
Excel workbook that mirrors the preview data shown in the app.
"""

import json
import sys
from pathlib import Path
from typing import List, Sequence, Tuple

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    print("ERROR: openpyxl is not installed. Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


def _select_columns(columns: Sequence[str]) -> List[Tuple[int, str]]:
    """Return (source_index, header) tuples for non-empty headers."""
    selected: List[Tuple[int, str]] = []
    for idx, name in enumerate(columns or []):
        if name is None:
            continue
        header = str(name).strip()
        if not header:
            continue
        selected.append((idx, header))
    return selected


def _append_rows(ws, column_map: Sequence[Tuple[int, str]], rows: Sequence[Sequence[object]]) -> None:
    for row in rows or []:
        values: List[object] = []
        for source_idx, _ in column_map:
            if source_idx < len(row):
                values.append(row[source_idx])
            else:
                values.append("")
        ws.append(values)


def build_workbook(output_path: Path, sheet_name: str, columns: List[str], rows: List[List[object]]) -> None:
    """
    Create a workbook with a single sheet, populate headers/data, and add an Excel Table.
    Only the columns present in the preview (non-empty headers) are included. Their
    original positions are preserved so, for example, removing column 3 still allows
    column 4 to appear as the third column in the table.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or "Sheet1"  # Excel sheet name limit per spec

    column_map = _select_columns(columns)
    if not column_map:
        column_map = [(0, "Column1")]

    # Write headers
    ws.append([header for _, header in column_map])

    # Write row data aligned with the selected column indices
    _append_rows(ws, column_map, rows)

    column_count = len(column_map)
    last_column_letter = get_column_letter(column_count)
    last_row_index = max(ws.max_row, 1)
    table_ref = f"A1:{last_column_letter}{last_row_index}"

    display_name = sheet_name.replace(" ", "_")[:255] or "Table1"
    table = Table(displayName=display_name, ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleMedium7",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    wb.save(output_path)


def main():
    if len(sys.argv) < 5:
        print("Usage: excelWriter.py <output_path> <sheet_name> <columns_json> <rows_json>", file=sys.stderr)
        sys.exit(1)

    output_path = Path(sys.argv[1])
    sheet_name = sys.argv[2] or "Sheet1"
    columns = json.loads(sys.argv[3]) if sys.argv[3] else []
    rows = json.loads(sys.argv[4]) if len(sys.argv) > 4 and sys.argv[4] else []

    build_workbook(output_path, sheet_name, columns, rows)


if __name__ == "__main__":
    main()