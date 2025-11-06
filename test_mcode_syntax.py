"""Test the actual M code syntax that would be generated"""
import json
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed")
    exit(1)

file_path = r"C:\Users\jpo64\Downloads\sample_powerquery_data.xlsx"
wb = openpyxl.load_workbook(file_path, read_only=False, data_only=True)

# Find Table1
for sheet in wb.worksheets:
    if hasattr(sheet, 'tables') and sheet.tables:
        for table_name, table_ref in sheet.tables.items():
            if table_name == "Table1":
                # Get table range
                if isinstance(table_ref, str):
                    ref = table_ref
                else:
                    ref = table_ref.ref
                table_range = sheet[ref]
                
                # Extract data
                rows = []
                for row in table_range:
                    rows.append([cell.value for cell in row])
                
                columns = [str(cell) if cell is not None else "" for cell in rows[0]]
                data_rows = rows[1:3]  # Just first 2 rows for testing
                
                # Generate M code exactly as the code does
                columns_m = "{" + ", ".join([f'"{col}"' for col in columns]) + "}"
                
                def format_value(val):
                    if val is None:
                        return "null"
                    elif isinstance(val, (int, float)):
                        return str(val)
                    elif isinstance(val, bool):
                        return "true" if val else "false"
                    else:
                        val_str = str(val)
                        val_str_escaped = val_str.replace('"', '""')
                        return f'"{val_str_escaped}"'
                
                rows_m = "{" + ", ".join([
                    "{" + ", ".join([
                        format_value(val)
                        for val in row[:len(columns)]
                    ]) + "}"
                    for row in data_rows
                ]) + "}"
                
                m_table = f"#table({columns_m}, {rows_m})"
                
                normalized_path = file_path.replace("\\", "/")
                source_description = f'Excel.Workbook(File.Contents("{normalized_path}"), null, true){{[Item="Table1",Kind="Table"]}}[Data]'
                
                source_step = f"Source = {m_table} // From: {source_description}"
                
                new_m_code = f"""let
    {source_step}
in
    Source"""
                
                print("Generated M code:")
                print("=" * 80)
                print(new_m_code)
                print("=" * 80)
                print(f"\nLine 5 (character 5-20): '{new_m_code.split(chr(10))[4][4:20] if len(new_m_code.split(chr(10))) > 4 else 'N/A'}'")
                print(f"\nTotal lines: {len(new_m_code.split(chr(10)))}")
                
                # Check for potential syntax issues
                lines = new_m_code.split("\n")
                for i, line in enumerate(lines, 1):
                    print(f"Line {i}: {repr(line)}")
                
                wb.close()
                break

